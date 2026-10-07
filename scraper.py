"""
PDPC Decisions Scraper
Extracts the Protection Obligation decisions from PDPC's enforcement decisions
(https://www.pdpc.gov.sg/organisations/regulations-decisions/enforcement-decisions)
and writes them to pdpc_decisions.xlsx.

The site is server-rendered: the listing page embeds an index of every decision, and each decision
page carries its date, summary and PDF. Plain requests + BeautifulSoup; no browser.
"""

import argparse
import io
import json
import re
import time
import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from pypdf import PdfReader

# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------

BASE_URL     = "https://www.pdpc.gov.sg"
LISTING_URL  = f"{BASE_URL}/organisations/regulations-decisions/enforcement-decisions?type=Commission%27s+Decisions"
OUTPUT_FILE  = "pdpc_decisions.xlsx"

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/122.0.0.0 Safari/537.36"
    ),
    "Accept-Language": "en-US,en;q=0.9",
}

DELAY_DETAIL = 1.5   # seconds between individual decision page requests

# ---------------------------------------------------------------------------
# Regex patterns
# ---------------------------------------------------------------------------

# Case citation: [YYYY] SGPDPC N or [YYYY] SGPDPCS N (newer summary format)
CITATION_RE = re.compile(
    r'\[\d{4}\]\s+SGPDPCS?\s+\d+(?:\s*\(NFA\))?',
    re.IGNORECASE,
)

# PDF link on detail pages
PDF_LINK_RE = re.compile(r'\.pdf', re.IGNORECASE)

# Context-aware financial penalty: "financial penalty of $X,XXX"
PENALTY_CONTEXT_RE = re.compile(
    r'(?:financial\s+penalty\s+of\s+|penalty\s+of\s+)'
    r'(?:S\$|\$)([\d,]+(?:\.\d{2})?)',
    re.IGNORECASE,
)

# Fallback: any standalone dollar amount
PENALTY_FALLBACK_RE = re.compile(
    r'(?:S\$|\$)([\d,]+(?:\.\d{2})?)',
    re.IGNORECASE,
)

# ---------------------------------------------------------------------------
# HTTP helpers (for detail pages)
# ---------------------------------------------------------------------------

def safe_get(session: requests.Session, url: str, retries: int = 3, backoff: float = 2.0):
    """GET url with retry + exponential backoff. Returns Response or None."""
    for attempt in range(retries):
        try:
            resp = session.get(url, timeout=20, headers=HEADERS)
            if resp.status_code == 429:
                wait = backoff * (2 ** attempt)
                print(f"  Rate limited. Waiting {wait:.0f}s...")
                time.sleep(wait)
                continue
            resp.raise_for_status()
            return resp
        except requests.exceptions.RequestException as e:
            wait = backoff * (2 ** attempt)
            print(f"  Request failed (attempt {attempt + 1}/{retries}): {e}")
            if attempt < retries - 1:
                time.sleep(wait)
    return None

# ---------------------------------------------------------------------------
# Extraction helpers
# ---------------------------------------------------------------------------

def extract_citation_from_pdf(session: requests.Session, pdf_url: str) -> str:
    """Download the decision PDF and search the first few pages for the case citation."""
    try:
        resp = session.get(pdf_url, timeout=30)
        resp.raise_for_status()
        reader = PdfReader(io.BytesIO(resp.content))
        if not reader.pages:
            return ""
        # Search first 3 pages — citation can appear on page 0 or 1
        for page in reader.pages[:3]:
            text = page.extract_text() or ""
            match = CITATION_RE.search(text)
            if match:
                return match.group(0).strip()
        return ""
    except Exception:
        return ""


def extract_citation(text: str) -> str:
    match = CITATION_RE.search(text)
    return match.group(0).strip() if match else ""


def extract_penalty(text: str) -> str:
    # Try context-aware extraction first
    match = PENALTY_CONTEXT_RE.search(text)
    if match:
        return f"${match.group(1)}"

    # If no dollar amount at all, determine whether there simply is no penalty
    has_dollar = PENALTY_FALLBACK_RE.search(text)
    if not has_dollar:
        lower = text.lower()
        no_penalty_keywords = [
            "no further action", "warning", "not in breach",
            "advisory notice", "directions were issued", "directions issued",
            "no breach",
        ]
        if any(kw in lower for kw in no_penalty_keywords):
            return "None"
        # No dollar amount and no recognisable decision keyword — leave blank
        return ""

    # Fallback: first dollar amount anywhere
    match = PENALTY_FALLBACK_RE.search(text)
    return f"${match.group(1)}" if match else ""


def normalise_case_name(title: str) -> str:
    """
    Convert 'Breach of the Protection Obligation by Acme Pte Ltd'
    → 'Re Acme Pte Ltd'
    Handles 'No Breach', 'No breach', 'Breach of Data Protection...' etc.
    """
    match = re.search(r'\bby\s+(.+)$', title, re.IGNORECASE)
    if match:
        return "Re " + match.group(1).strip()
    return title


def extract_description(rte_soup) -> str:
    """Return the first substantive paragraph from the .rte block."""
    if rte_soup is None:
        return ""
    for p in rte_soup.find_all("p"):
        text = p.get_text(" ", strip=True)
        if len(text) > 80 and not CITATION_RE.match(text):
            return text
    return rte_soup.get_text(" ", strip=True)[:600]

# ---------------------------------------------------------------------------
# Phase 1: the listing. The enforcement decisions page is server-rendered (Next.js) and embeds an index of
# every decision (title and link) in its page data, so one request lists them all; no browser needed.
# ---------------------------------------------------------------------------

INDEX_ITEM_RE = re.compile(
    r'\{"label":"((?:[^"\\]|\\.)*)","url":"(/organisations/regulations-decisions/enforcement-decisions/[^"]+)",'
    r'"isListingDetail":true'
)

OBLIGATIONS = [
    "Protection", "Accountability", "Consent", "Notification", "Purpose Limitation", "Openness", "Access",
    "Correction", "Accuracy", "Retention Limitation", "Transfer Limitation", "Data Breach Notification",
]


def page_data(html: str) -> str:
    """The Next.js flight data embedded in a page (self.__next_f.push chunks), decoded."""
    chunks = re.findall(r'self\.__next_f\.push\(\[1,"(.*?)"\]\)', html, re.S)
    return "".join(json.loads('"' + c + '"') for c in chunks)


def resolve_reference(data: str, value: str) -> str:
    """Next.js sends long strings as separate text rows ("27:T<hex byte length>,<text>") and refers to them as
    "$27"; return the text a reference points to, or the value unchanged."""
    ref = re.fullmatch(r"\$([0-9a-f]+)", value)
    if not ref:
        return value
    row = re.search(rf"(?:^|\n){ref.group(1)}:T([0-9a-f]+),", data)
    if not row:
        return value
    return data[row.end():].encode("utf-8")[: int(row.group(1), 16)].decode("utf-8", "ignore")


def obligations_in(title: str) -> list[str]:
    """Obligations named in a decision title, e.g. 'Breach of the Accountability and Protection Obligations'."""
    found = [o for o in OBLIGATIONS if re.search(rf"\b{o}\b", title, re.IGNORECASE)]
    # "Data Breach Notification" also contains "Notification"
    if "Data Breach Notification" in found and not re.search(r"(?<!Breach )\bNotification\b", title):
        found.remove("Notification")
    return found


def fetch_listings(session: requests.Session) -> list[dict]:
    print(f"Phase 1: Reading the decision index from {LISTING_URL} ...")
    resp = safe_get(session, LISTING_URL)
    if resp is None:
        return []
    items, seen = [], set()
    for m in INDEX_ITEM_RE.finditer(page_data(resp.text)):
        url = m.group(2)
        if url in seen:
            continue
        seen.add(url)
        items.append({"title": json.loads('"' + m.group(1) + '"'), "url": url})
    print(f"  {len(items)} decisions in the index")
    return items

# ---------------------------------------------------------------------------
# Phase 2: each decision page: published date, summary (with the penalty) and the decision PDF
# ---------------------------------------------------------------------------

CONTENT_RE = re.compile(r'"data":\{"content":"((?:[^"\\]|\\.)*)"\}')


def decision_types(summary: str) -> str:
    lower = summary.lower()
    kinds = []
    if re.search(r"no breach|not (?:to be )?in breach|did not breach", lower):
        kinds.append("Not in Breach")
    if re.search(r"financial penalt(?:y|ies)|penalt(?:y|ies) of", lower):
        kinds.append("Financial Penalty")
    if re.search(r"\bdirections?\b", lower):
        kinds.append("Directions")
    if "warning" in lower:
        kinds.append("Warning")
    if "undertaking" in lower:
        kinds.append("Undertaking")
    return ", ".join(kinds)


def fetch_detail(session: requests.Session, item: dict, read_pdf: bool = True) -> dict | None:
    """One decision as a row, or None if the page can't be read."""
    full_url = BASE_URL + item["url"]
    resp = safe_get(session, full_url)
    if resp is None:
        return None
    soup = BeautifulSoup(resp.text, "html.parser")
    data = page_data(resp.text)

    date_el = soup.select_one(".page-banner__date")
    date = re.sub(r"^\s*Published on\s*", "", date_el.get_text(" ", strip=True)) if date_el else ""

    match = CONTENT_RE.search(data)
    content = BeautifulSoup(resolve_reference(data, json.loads('"' + match.group(1) + '"')), "html.parser") if match else None
    summary = re.sub(r"\s*Click here to find out more\.?\s*$", "", content.get_text(" ", strip=True)) if content else ""

    pdf_url = ""
    if content:
        for a in content.find_all("a", href=True):
            href = a["href"]
            if href.startswith("/assets/") or PDF_LINK_RE.search(href):
                pdf_url = BASE_URL + href if href.startswith("/") else href
                break

    title = item["title"]
    citation = extract_citation(summary)
    if not citation and pdf_url and read_pdf:
        citation = extract_citation_from_pdf(session, pdf_url)

    return {
        "case_name":     normalise_case_name(title),
        "citation":      citation,
        "date":          date,
        "obligations":   ", ".join(obligations_in(title)),
        "decision_type": decision_types(summary),
        "description":   summary,
        "penalty":       extract_penalty(summary),
        "url":           full_url,
        "title":         title,
        "summary_lower": summary.lower(),
    }


def is_protection_case(row: dict) -> bool:
    """Protection Obligation cases: named in the title, or (older 'Data Protection Provisions' titles and
    undertakings) described in the summary."""
    if "Protection" in row["obligations"].split(", "):
        return True
    return bool(re.search(r"protection obligation|section 24\b", row["summary_lower"]))

# ---------------------------------------------------------------------------
# Phase 3: Write Excel
# ---------------------------------------------------------------------------

HEADER = [
    "Case Name",
    "Case Citation",
    "Date",
    "Obligations",
    "Decision Type",
    "Case Description",
    "Financial Penalty",
    "URL",
]

COL_WIDTHS = [60, 22, 14, 35, 30, 80, 20, 80]


def write_excel(rows: list[dict], filename: str = OUTPUT_FILE) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "PDPC Decisions"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E79")
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    wrap_align   = Alignment(vertical="top", wrap_text=True)

    for col_idx, header_text in enumerate(HEADER, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header_text)
        cell.font      = header_font
        cell.fill      = header_fill
        cell.alignment = center_align

    col_letters = ["A", "B", "C", "D", "E", "F", "G", "H"]
    for letter, width in zip(col_letters, COL_WIDTHS):
        ws.column_dimensions[letter].width = width

    ws.freeze_panes = "A2"

    for row_idx, row in enumerate(rows, start=2):
        values = [
            row["case_name"],
            row["citation"],
            row["date"],
            row["obligations"],
            row["decision_type"],
            row["description"],
            row["penalty"],
            row["url"],
        ]
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = wrap_align

    wb.save(filename)
    print(f"\nSaved {len(rows)} rows to {filename}")

# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(description="Scrape PDPC Protection Obligation decisions into Excel.")
    parser.add_argument("--output", default=OUTPUT_FILE, help=f"Excel file to write (default: {OUTPUT_FILE}).")
    parser.add_argument("--limit", type=int, default=0, help="Stop after this many decision pages (for a quick check).")
    parser.add_argument("--delay", type=float, default=DELAY_DETAIL, help="Seconds between decision pages.")
    parser.add_argument("--no-pdf", action="store_true", help="Skip downloading PDFs to find citations.")
    args = parser.parse_args()

    session = requests.Session()
    listings = fetch_listings(session)
    if not listings:
        print("No listings found. Exiting.")
        return
    if args.limit:
        listings = listings[:args.limit]

    # Phase 2: every decision page; keep the Protection Obligation cases
    print(f"\nPhase 2: Reading {len(listings)} decision pages...")
    rows = []
    for i, item in enumerate(listings, start=1):
        row = fetch_detail(session, item, read_pdf=not args.no_pdf)
        if row and is_protection_case(row):
            if not row["obligations"]:
                row["obligations"] = "Protection"  # older titles say "Data Protection Provisions"
            rows.append(row)
            print(f"  [{i:3d}/{len(listings)}] {row['title'][:70]}")
        time.sleep(args.delay)
    print(f"  {len(rows)} Protection Obligation decisions")

    # Phase 3: write Excel
    print("\nPhase 3: Writing Excel file...")
    write_excel(rows, args.output)
    print("Done.")


if __name__ == "__main__":
    main()
